from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables" / "晚年IP文书包-v4"
OUT.mkdir(parents=True, exist_ok=True)

FONT = "STSong"
HEAD_FONT = "STHeiti"
TIP_GRAY = "595959"
TABLE_WIDTH = 8966
TABLE_INDENT = -16
FIELD = "______________"
LONG_FIELD = "________________________"
TABLE_FONT_SIZE = 10
TABLE_HEADER_SIZE = 10.5
TABLE_LINE_PT = 15


def set_font(run, name=FONT, size=14, bold=False, underline=False, color=None, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.underline = underline
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def paragraph_format(p, line_pt=24, align=None, before=0, after=0, first_line=0, left=0):
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing = Pt(line_pt)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if first_line:
        p.paragraph_format.first_line_indent = Inches(first_line)
    if left:
        p.paragraph_format.left_indent = Inches(left)


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.268)
    section.page_height = Inches(11.693)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.25)
    section.right_margin = Inches(1.25)
    section.header_distance = Inches(0.5)
    section.footer_distance = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)

    heading1 = doc.styles["Heading 1"]
    heading1.font.name = HEAD_FONT
    heading1._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
    heading1.font.size = Pt(22)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(0, 0, 0)
    heading1.paragraph_format.space_before = Pt(17)
    heading1.paragraph_format.space_after = Pt(16.5)
    heading1.paragraph_format.line_spacing = 2.4

    heading2 = doc.styles["Heading 2"]
    heading2.font.name = HEAD_FONT
    heading2._element.rPr.rFonts.set(qn("w:eastAsia"), HEAD_FONT)
    heading2.font.size = Pt(16)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(0, 0, 0)
    heading2.paragraph_format.space_before = Pt(13)
    heading2.paragraph_format.space_after = Pt(13)
    heading2.paragraph_format.line_spacing = 1.72

    for section in doc.sections:
        section.header.paragraphs[0].text = ""
        section.footer.paragraphs[0].text = ""


def add_ref_paragraph(
    doc,
    text="",
    size=14,
    bold=False,
    color=None,
    align=WD_ALIGN_PARAGRAPH.LEFT,
    line_pt=24,
    before=0,
    after=0,
    italic=False,
):
    p = doc.add_paragraph()
    paragraph_format(p, line_pt=line_pt, align=align, before=before, after=after)
    if text:
        r = p.add_run(text)
        set_font(r, FONT, size, bold, False, color, italic)
    return p


def add_rich_paragraph(doc, parts, size=14, align=WD_ALIGN_PARAGRAPH.LEFT, line_pt=24, before=0, after=0):
    p = doc.add_paragraph()
    paragraph_format(p, line_pt=line_pt, align=align, before=before, after=after)
    for part in parts:
        if isinstance(part, str):
            text, bold, underline, color, italic, font = part, False, False, None, False, FONT
        else:
            text = part.get("text", "")
            bold = part.get("bold", False)
            underline = part.get("underline", False)
            color = part.get("color")
            italic = part.get("italic", False)
            font = part.get("font", FONT)
        run = p.add_run(text)
        set_font(run, font, size, bold, underline, color, italic)
    return p


def add_main_title(doc, text):
    return add_ref_paragraph(
        doc,
        text,
        size=36,
        bold=False,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        line_pt=27,
        after=0,
    )


def add_large_center(doc, text, size=26, bold=True):
    return add_ref_paragraph(
        doc,
        text,
        size=size,
        bold=bold,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        line_pt=27,
        after=0,
    )


def add_heading(doc, text, level=1, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph(style="Heading 1" if level == 1 else "Heading 2")
    size = 22 if level == 1 else 16
    paragraph_format(p, line_pt=24, align=align, before=8 if level == 1 else 5, after=5)
    r = p.add_run(text)
    set_font(r, HEAD_FONT, size, True, color="000000")
    return p


def add_tip(doc, text, label="提示："):
    return add_rich_paragraph(
        doc,
        [
            {"text": label, "bold": True, "color": TIP_GRAY},
            {"text": text, "color": TIP_GRAY},
        ],
        size=14,
        line_pt=24,
        after=0,
    )


def add_plain_explanation(doc, text):
    return add_ref_paragraph(doc, text, size=14, line_pt=24, after=0)


def add_blank(doc, label="", chars=34, lines=1):
    if label:
        add_ref_paragraph(doc, label + "：" + "_" * chars, size=14, line_pt=24)
    else:
        add_ref_paragraph(doc, "_" * chars, size=14, line_pt=24)
    for _ in range(lines - 1):
        add_ref_paragraph(doc, "_" * chars, size=14, line_pt=24)


def add_signature(doc, role="本人"):
    add_ref_paragraph(
        doc,
        f"{role}签字（按手印）：________________________    日期：____年__月__日",
        size=14,
        line_pt=24,
    )


def set_table_borders(table, color="000000", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths):
    if sum(widths) != TABLE_WIDTH:
        widths[-1] += TABLE_WIDTH - sum(widths)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    # Long Chinese labels and underline blanks make Word's auto-fit algorithm
    # collapse narrow columns (the symptom seen in the user's screenshots).
    # Keep the reference table width, but make the declared grid authoritative.
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT))
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")
    set_table_borders(table)

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[index]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def mark_row_no_split(row):
    """Keep a table row intact when Word moves it to the next page."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def mark_header_row(row):
    """Repeat the first table row when a table continues on a new page."""
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        tr_pr.append(OxmlElement("w:tblHeader"))


def set_cell_margins(cell, top=55, start=70, bottom=55, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + name))
        if node is None:
            node = OxmlElement("w:" + name)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_clear_shading(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "auto")


def compact_table_text(value):
    text = str(value)
    replacements = {
        "____________________________________________________________": "________________________",
        "____________________________": "______________",
        "________________________": "______________",
        "________________": "__________",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def ref_table(doc, headers, rows, widths, font_size=TABLE_FONT_SIZE):
    font_size = min(font_size, TABLE_FONT_SIZE)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, list(widths))
    mark_header_row(table.rows[0])
    mark_row_no_split(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_clear_shading(cell)
        p = cell.paragraphs[0]
        paragraph_format(p, line_pt=TABLE_LINE_PT, align=WD_ALIGN_PARAGRAPH.CENTER)
        r = p.add_run(compact_table_text(header))
        set_font(r, FONT, TABLE_HEADER_SIZE, True)

    for row_values in rows:
        cells = table.add_row().cells
        mark_row_no_split(table.rows[-1])
        for i, value in enumerate(row_values):
            cell = cells[i]
            cell.text = ""
            set_clear_shading(cell)
            p = cell.paragraphs[0]
            paragraph_format(
                p,
                line_pt=TABLE_LINE_PT,
                align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT,
            )
            lines = compact_table_text(value).split("\n")
            for line_index, line in enumerate(lines):
                if line_index:
                    p.add_run().add_break()
                r = p.add_run(line)
                set_font(r, FONT, font_size)
    # Re-apply geometry after adding rows so every row carries the same tcW
    # values instead of relying on Word's inherited defaults.
    set_table_geometry(table, list(widths))
    add_ref_paragraph(doc, "", size=14, line_pt=8)
    return table


def add_table_intro(doc, text):
    return None


def add_after_table(doc, text):
    return None


def add_choice_block(doc, title, explanation, choices, field_label="我还想说明"):
    add_heading(doc, title, level=2)
    if explanation:
        add_plain_explanation(doc, explanation)
    for choice in choices:
        add_ref_paragraph(doc, "□ " + choice, size=14, line_pt=24)
    add_blank(doc, field_label, chars=58, lines=2)


def add_witness(doc, index):
    add_heading(doc, f"见证人{index}签字页", level=2)
    add_blank(doc, f"见证人{index}姓名", 26)
    add_blank(doc, f"身份证号", 28)
    add_blank(doc, f"与本人关系", 26)
    add_blank(doc, f"联系电话", 26)
    add_ref_paragraph(
        doc,
        "本人于____年__月__日在________________见到本人签字/按手印。本人能说出大概意思，签字时未见明显欺骗、胁迫或强迫。",
        size=14,
        line_pt=24,
    )
    add_signature(doc, f"见证人{index}")


def add_storage_table(doc):
    add_heading(doc, "八、签完以后怎么保存", level=1)
    add_heading(doc, "（一）原件、复印件和电子版放在哪里", level=2)
    add_ref_paragraph(doc, "联系人找不到文件，就等于没填。原件、复印件、电子版都要有人知道在哪里。", size=14, line_pt=24)
    ref_table(
        doc,
        ["资料", "放在哪里/交给谁", "什么时候更新"],
        [
            ["原件", "____________________________", "____年__月__日"],
            ["复印件一", "____________________________", "____年__月__日"],
            ["复印件二", "____________________________", "____年__月__日"],
            ["电子版", "____________________________", "____年__月__日"],
            ["药盒/病历/检查报告", "____________________________", "____年__月__日"],
            ["是否告知常去医院", "□已告知________  □还没有", "____年__月__日"],
        ],
        [1900, 5100, 1966],
        font_size=14,
    )
    add_heading(doc, "（二）修改、撤回和版本记录", level=2)
    add_ref_paragraph(doc, "不要在旧版本上乱改。重新打印、重新签字后，旧版本第一页写“已撤回/不再使用”。", size=14, line_pt=24)
    ref_table(
        doc,
        ["日期/版本", "改动/撤回", "新版本告知谁", "本人签字"],
        [
            ["____年__月__日\n版本____", "□新 □改 □撤旧\n________", "________", "________"],
            ["____年__月__日\n版本____", "□新 □改 □撤旧\n________", "________", "________"],
            ["____年__月__日\n版本____", "□新 □改 □撤旧\n________", "________", "________"],
            ["____年__月__日\n版本____", "□新 □改 □撤旧\n________", "________", "________"],
        ],
        [1800, 3300, 2300, 1566],
        font_size=14,
    )


def build_health_package(doc):
    add_main_title(doc, "生前预嘱与健康照护包")
    add_large_center(doc, "先从急用信息开始", size=26, bold=True)
    for item in [
        "□ 今天时间少：先填联系人、常用药、过敏、常去医院。",
        "□ 拿不准：空着或写“不知道”，不要替医生下结论。",
        "□ 不写银行卡密码、验证码、支付密码。",
    ]:
        add_ref_paragraph(doc, item, size=14, line_pt=24)

    add_heading(doc, "一、遇到事情，先看哪一页", level=1)
    ref_table(
        doc,
        ["遇到的情况", "先看哪份", "马上做什么"],
        [
            ["急诊、住院", "健康说明书", "先给医生看病、药、过敏、联系人"],
            ["请家人/护工照顾", "生活照护意愿书", "按本人习惯安排吃饭、洗澡、探视"],
            ["本人说不清话", "我的五个愿望", "先问医生病情，再看本人写过的想法"],
            ["签字当天", "本人签字和见证", "确认本人自愿，见证人在场签字"],
            ["以后找文件/改版本", "保存和更新记录", "找到最新版，避免拿错旧版本"],
        ],
        [2400, 2450, 4116],
        font_size=14,
    )

    add_heading(doc, "二、本人怎么填", level=1)
    ref_table(
        doc,
        ["顺序", "先做什么", "填到什么程度", "完成"],
        [
            ["1", "先填联系人、常用药、过敏、常去医院", "今天只能填一页，就先填这些", "□"],
            ["2", "照药盒、病历、检查单补病名和药名", "不会写就抄；不知道就写“不知道”", "□"],
            ["3", "写生活习惯和照护偏好", "让家人、护工看得懂就行", "□"],
            ["4", "和家人谈医疗愿望", "不懂的选“问医生”，不要硬选", "□"],
            ["5", "本人签字，告诉联系人文件放哪里", "联系人能找到最新版", "□"],
        ],
        [850, 3900, 3216, 1000],
        font_size=14,
    )

    add_heading(doc, "三、家人拿到后怎么用", level=1)
    for item in [
        "□ 本人还能清楚说话时，先问本人；这份文件只帮助家人少漏事。",
        "□ 急诊、住院时，先看健康说明书，不要只凭家人回忆。",
        "□ 遇到抢救、ICU、呼吸机等决定，先让医生说明病情，再看“我的五个愿望”。",
        "□ 找不到原件时，先联系保存记录里的第一联系人；不要直接拿旧版本用。",
    ]:
        add_ref_paragraph(doc, item, size=14, line_pt=24)

    doc.add_page_break()
    add_heading(doc, "四、个人健康说明书", level=1)
    add_ref_paragraph(doc, "给急诊、住院、家属核对用。这里只写事实；治疗怎么做，听医生当时说明。", size=14, line_pt=24)

    add_heading(doc, "（一）我是谁，平时能不能自己照顾自己", level=2)
    ref_table(
        doc,
        ["项目", "请填写"],
        [
            ["姓名/曾用名", FIELD],
            ["性别/出生日期", FIELD],
            ["身份证号", FIELD],
            ["常住地址", FIELD],
            ["血型（不知道可空）", FIELD],
            ["走路（近三个月）", "□自己走  □需要扶  □不能自己走"],
            ["吃饭（近三个月）", "□自己吃  □需要提醒  □需要喂"],
            ["说话/听话（近三个月）", "□正常  □听不清  □说不清  □需要慢慢说"],
        ],
        [2600, 6366],
        font_size=14,
    )

    add_heading(doc, "（二）我有哪些长期需要关注的疾病", level=2)
    add_ref_paragraph(doc, "长期吃药、长期复查的病，都写。病名照病历写。", size=14, line_pt=24)
    ref_table(
        doc,
        ["疾病", "发现时间/医院", "注意事项"],
        [
            ["1. ____________", "____________________________", "____________________________"],
            ["2. ____________", "____________________________", "____________________________"],
            ["3. ____________", "____________________________", "____________________________"],
            ["4. ____________", "____________________________", "____________________________"],
        ],
        [2400, 3150, 3416],
        font_size=14,
    )
    add_blank(doc, "最近一次检查报告放在", chars=42)

    add_heading(doc, "（三）我每天吃什么药", level=2)
    add_ref_paragraph(doc, "药名照药盒写全名。备用药也写。", size=14, line_pt=24)
    ref_table(
        doc,
        ["药名", "用处", "用法", "位置"],
        [
            ["________________", "________________", "________________", "________________"],
            ["________________", "________________", "________________", "________________"],
            ["________________", "________________", "________________", "________________"],
            ["________________", "________________", "________________", "________________"],
            ["备用/急救药：________", "________________", "________________", "________________"],
        ],
        [2500, 2350, 2350, 1766],
        font_size=14,
    )
    add_blank(doc, "药盒和说明书放在", chars=42)

    add_heading(doc, "（四）我有没有过敏", level=2)
    add_ref_paragraph(doc, "记不清，就写“不知道”。", size=14, line_pt=24)
    add_ref_paragraph(doc, "药物过敏：□没有    □有，哪种药：________________，出现什么反应：________________")
    add_ref_paragraph(doc, "食物过敏：□没有    □有，哪种食物：____________，出现什么反应：________________")
    add_ref_paragraph(doc, "其他过敏：□没有    □有，什么东西：____________，出现什么反应：________________")

    add_heading(doc, "（五）我平时去哪里看病", level=2)
    ref_table(
        doc,
        ["项目", "请填写"],
        [
            ["常去医院/社区医院", LONG_FIELD],
            ["科室/医生", LONG_FIELD],
            ["医院/医生电话", LONG_FIELD],
            ["病历/报告位置", LONG_FIELD],
            ["医保/保险资料", LONG_FIELD],
        ],
        [2500, 6466],
        font_size=14,
    )

    add_heading(doc, "（六）第一联系人和第二联系人", level=2)
    ref_table(
        doc,
        ["顺序", "姓名/关系", "电话", "可帮事项"],
        [
            ["第一联系人", "________________", "________________", "________________"],
            ["第二联系人", "________________", "________________", "________________"],
            ["其他联系人", "________________", "________________", "________________"],
        ],
        [1500, 2650, 2000, 2816],
        font_size=14,
    )

    add_heading(doc, "（七）我平时需要别人注意什么", level=2)
    add_blank(doc, "照护提醒（跌倒、怕冷、忌口、洗澡、夜间等）", chars=44, lines=4)
    add_blank(doc, "最近一次更新日期", chars=28)

    add_heading(doc, "五、个人生活照护意愿书", level=1)
    add_ref_paragraph(doc, "给家人、护工、养老机构看。写日常习惯，不写谁出钱、谁继承。", size=14, line_pt=24)

    add_choice_block(
        doc,
        "（一）如果需要长期照护，我更愿意住在哪里",
        "按顺序写 1、2、3。拿不准，就写“再商量”。",
        [
            "尽可能住在自己家里",
            "住在子女/亲属家里",
            "住养老院或护理机构",
            "住医院或舒缓照护机构",
            "和家人再商量",
        ],
        field_label="我最在意的居住条件",
    )

    add_choice_block(
        doc,
        "（二）我喜欢怎样吃饭",
        "",
        [
            "口味：□清淡  □偏咸  □偏甜  □其他________________",
            "食物：□软一点  □切小块  □流质/半流质  □普通饮食",
            "我不吃/不能吃：____________________________",
            "吃饭时间：早____点；午____点；晚____点",
        ],
        field_label="我最喜欢的一样食物",
    )

    add_choice_block(
        doc,
        "（三）我喜欢怎样睡觉、洗澡和穿衣",
        "",
        [
            "起床时间：____点左右；睡觉时间：____点左右",
            "午睡：□需要  □不需要，大约____分钟",
            "洗澡/擦身：□淋浴  □盆浴  □擦身；需要注意________________",
            "穿衣：□自己选  □家人帮选；我喜欢的衣服/颜色________________",
            "如厕和隐私：进入房间前请先________________",
        ],
        field_label="我最在意的一件小事",
    )

    add_choice_block(
        doc,
        "（四）我希望谁来看我、怎样陪我",
        "",
        [
            "探视：□欢迎  □请先打电话  □每天不超过____人",
            "我希望常来陪我的人：____________________________",
            "我暂时不想见的人：____________________________",
            "我喜欢：□聊天  □听音乐  □看照片  □读报  □晒太阳  □散步",
            "我希望别人不要：____________________________",
        ],
        field_label="我最需要的陪伴",
    )

    add_choice_block(
        doc,
        "（五）我对隐私、尊严和信仰的愿望",
        "",
        [
            "请在做检查、换衣服、洗澡前先告诉我",
            "不要当着很多人的面谈我的病",
            "我希望听到的音乐/看到的照片：________________",
            "我的宗教或精神习惯：□没有特别要求  □请照顾到：________________",
            "我最不希望别人：________________",
        ],
        field_label="我最在意的一件事",
    )

    add_heading(doc, "六、我的五个愿望", level=1)
    add_ref_paragraph(doc, "给家人和医生沟通用。不是医嘱；具体病情要先听医生解释，再看本人写过的想法。", size=14, line_pt=24)

    add_heading(doc, "第一个愿望：我希望接受什么医疗服务", level=2)
    ref_table(
        doc,
        ["事项", "我的选择", "想先知道"],
        [
            ["止痛/缓解不适", "□优先  □视情况  □问医生", "__________"],
            ["必要检查", "□接受  □少做  □问医生", "__________"],
            ["抢救/复苏", "□接受  □不接受  □问医生", "__________"],
            ["呼吸机/插管", "□接受  □不接受  □问医生", "__________"],
            ["ICU/长期住院", "□接受  □不接受  □问医生", "__________"],
            ["输血/透析/抗感染", "□接受  □不接受  □问医生", "__________"],
            ["其他：________", "□接受  □不接受  □问医生", "__________"],
        ],
        [2450, 4050, 2466],
        font_size=14,
    )

    add_heading(doc, "第二个愿望：如果治疗只能维持身体，我希望怎么办", level=2)
    add_ref_paragraph(doc, "最终听医生当时说明。", size=14, line_pt=24)
    ref_table(
        doc,
        ["病情", "呼吸机/复苏", "营养/补液"],
        [
            ["生命末期", "□用  □不用  □舒缓  □问医生", "□用  □不用  □舒缓  □问医生"],
            ["不可逆昏迷", "□用  □不用  □舒缓  □问医生", "□用  □不用  □舒缓  □问医生"],
            ["植物状态", "□用  □不用  □舒缓  □问医生", "□用  □不用  □舒缓  □问医生"],
            ["其他：________", "□用  □不用  □舒缓  □问医生", "□用  □不用  □舒缓  □问医生"],
        ],
        [2000, 3483, 3483],
        font_size=14,
    )
    add_ref_paragraph(doc, "不管选什么，都请尽量做到：□止痛  □缓解喘  □清洁  □陪伴  □安静  □其他________________")
    add_blank(doc, "我希望医生重点说明", chars=64, lines=3)

    add_heading(doc, "第三个愿望：我希望别人怎样对待我", level=2)
    for item in [
        "□ 保持身体清洁，做检查、换衣服和洗澡前先告诉我。",
        "□ 让我尽量有隐私，不在陌生人面前谈我的病情。",
        "□ 身体允许时，让我听音乐、看照片或有人陪。",
        "□ 尽量让我在熟悉的地方。",
        "□ 请记得我的饮食习惯、信仰和不想见的人。",
        "□ 不要吓我、骂我、骗我，逼我作决定。",
    ]:
        add_ref_paragraph(doc, item, size=14, line_pt=24)
    add_blank(doc, "我还想说的一件事", chars=64, lines=3)

    add_heading(doc, "第四个愿望：我希望家人和朋友知道什么", level=2)
    add_blank(doc, "我想对家人说的话", chars=64, lines=4)
    ref_table(
        doc,
        ["顺序", "通知谁/电话", "要告诉他/她"],
        [
            ["第一位", "________________", "________________"],
            ["第二位", "________________", "________________"],
            ["第三位", "________________", "________________"],
            ["家里哪些事不要争", "________________", "________________"],
            ["告别/纪念的愿望", "________________", "________________"],
        ],
        [1800, 3000, 4166],
        font_size=14,
    )
    add_ref_paragraph(doc, "器官捐献、遗体安排、保险受益人，请另办正式手续。", size=14, line_pt=24)

    add_heading(doc, "第五个愿望：我不能表达时，谁帮助我把话告诉医生", level=2)
    add_ref_paragraph(doc, "此处只用于医疗和照护沟通，不授权处理财产。", size=14, line_pt=24)
    ref_table(
        doc,
        ["顺序", "姓名/关系", "电话", "信任原因", "不能做什么"],
        [
            ["第一人", "________________", "________________", "________________", "________________"],
            ["第二人", "________________", "________________", "________________", "________________"],
            ["第三人/备选", "________________", "________________", "________________", "________________"],
        ],
        [1000, 2000, 1650, 2450, 1866],
        font_size=14,
    )
    add_blank(doc, "如果家人意见不一样，我希望先这样处理", chars=64, lines=3)

    add_heading(doc, "七、本人签字和见证", level=1)
    add_ref_paragraph(doc, "本人确认后签字，不要代签。", size=14, line_pt=24)
    add_signature(doc, "本人")
    add_witness(doc, 1)
    add_witness(doc, 2)
    add_ref_paragraph(doc, "以后改变想法，就重新填、重新签，并告诉家人旧版本不用了。", size=14, line_pt=24)

    add_storage_table(doc)
    add_heading(doc, "九、家属沟通记录", level=1)
    add_ref_paragraph(doc, "记录谁已经知道文件位置和本人想法。以后住院、请护工、换联系人时，先看这里。", size=14, line_pt=24)
    ref_table(
        doc,
        ["日期/地点", "参加人", "说清楚了什么", "下一步"],
        [
            ["____年__月__日\n________________", "________________", "________________", "________________"],
            ["____年__月__日\n________________", "________________", "________________", "________________"],
            ["____年__月__日\n________________", "________________", "________________", "________________"],
            ["____年__月__日\n________________", "________________", "________________", "________________"],
        ],
        [1900, 2100, 2700, 2266],
        font_size=14,
    )

    add_heading(doc, "十、填写核对", level=1)
    for item in [
        "□ 我写了自己的姓名、身份证号、常用医院和第一/第二联系人。",
        "□ 我写了慢性病、常用药和过敏；不知道的地方写了“不知道”。",
        "□ 我写了吃饭、睡觉、洗澡、居住、探视和隐私方面的愿望。",
        "□ 医疗选择里，不懂的地方写了“问医生”。",
        "□ 第一联系人知道这份文件放在哪里，也愿意在需要时帮我沟通。",
        "□ 本人已签字；见证人已签字。",
        "□ 我记录了原件、复印件、电子版的位置和下一次更新日期。",
        "□ 健康、家庭关系或想法变化后，我会重新检查这份文件。",
    ]:
        add_ref_paragraph(doc, item, size=14, line_pt=24)
    add_blank(doc, "本次填写日期", chars=28)
    add_blank(doc, "下一次复核日期", chars=28)
    add_signature(doc, "本人")

    add_heading(doc, "十一、参考来源", level=1)
    add_ref_paragraph(doc, "北京生前预嘱推广协会：www.lwpa.org.cn", size=14, line_pt=24)
    add_ref_paragraph(doc, "“我的五个愿望”登记平台：www.livingwill.org.cn", size=14, line_pt=24)

    add_heading(doc, "十二、补充记录", level=1)
    add_ref_paragraph(doc, "不要写银行卡密码、验证码。", size=14, line_pt=24)
    add_blank(doc, "补充记录一", chars=64, lines=2)
    add_blank(doc, "补充记录二", chars=64, lines=2)
    add_blank(doc, "补充记录三", chars=64, lines=2)


def main():
    path = OUT / "01-生前预嘱与健康照护包-v4.docx"
    doc = Document()
    style_doc(doc)
    build_health_package(doc)
    doc.save(path)
    print(path)


if __name__ == "__main__":
    main()
